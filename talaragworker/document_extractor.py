from __future__ import annotations

from io import BytesIO


def extract_text(filename: str | None, content_type: str | None, content_bytes: bytes) -> str:
    lowered_name = (filename or "").lower()
    lowered_type = (content_type or "").lower()

    if lowered_name.endswith(".txt") or lowered_type.startswith("text/plain"):
        return _extract_txt(content_bytes)
    if lowered_name.endswith(".pdf") or lowered_type == "application/pdf":
        return _extract_pdf(content_bytes)
    if lowered_name.endswith(".xlsx") or "spreadsheetml" in lowered_type:
        return _extract_xlsx(content_bytes)
    if lowered_name.endswith(".pptx") or "presentationml" in lowered_type:
        return _extract_pptx(content_bytes)
    return _extract_txt(content_bytes)


def _extract_txt(content_bytes: bytes) -> str:
    if not content_bytes:
        return ""

    for encoding in ("utf-8", "latin-1"):
        try:
            return content_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content_bytes.decode("utf-8", errors="ignore")


def _extract_pdf(content_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""

    reader = PdfReader(BytesIO(content_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(part.strip() for part in pages if part.strip())


def _extract_xlsx(content_bytes: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""

    workbook = load_workbook(BytesIO(content_bytes), read_only=True, data_only=True)
    rows: list[str] = []
    for worksheet in workbook.worksheets:
        rows.append(f"# Sheet: {worksheet.title}")
        for values in worksheet.iter_rows(values_only=True):
            cells = [str(value).strip() for value in values if value not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_pptx(content_bytes: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return ""

    presentation = Presentation(BytesIO(content_bytes))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"# Slide {index}\n" + "\n".join(texts))
    return "\n\n".join(slides)
